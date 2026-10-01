"""Render report models into native, editable PowerPoint presentations."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from wood_reports.model import (
    ChartReference,
    Finding,
    Narrative,
    PublicationTable,
    Report,
    Section,
    resolve_chart_artifact,
)
from wood_reports.primitives import list_entries, text_runs
from wood_reports.theme import PublicationTheme

_DEFAULT_TABLE_MAX_ROWS = 8


class PowerPointRenderError(ValueError):
    """A rendering failure with the report element responsible for it."""

    def __init__(self, element: str, message: str) -> None:
        self.element = element
        super().__init__(f"{element}: {message}")


class PowerPointRenderer:
    """Render a report independently once its logical charts are resolved."""

    def render(self, report: Report, destination: Path, *, artifact_root: Path) -> Path:
        return _PowerPointDocument(report.theme).render(
            report, destination, artifact_root=artifact_root
        )


class _PowerPointDocument:
    def __init__(self, theme: PublicationTheme) -> None:
        self.theme = theme
        self._slide_width = Inches(theme.geometry.slide_width)
        self._slide_height = Inches(theme.geometry.slide_height)
        self._margin = Inches(theme.geometry.slide_margin)
        self._content_top = Inches(theme.geometry.content_top)
        self._content_height = Inches(theme.geometry.content_height)
        self._primary = self._color(theme.colors.primary)
        self._muted = self._color(theme.colors.text_muted)
        self._content_width = self._slide_width - 2 * self._margin

    @staticmethod
    def _color(value: str) -> Any:
        return RGBColor.from_string(value.removeprefix("#"))  # type: ignore[no-untyped-call]

    def render(self, report: Report, destination: Path, *, artifact_root: Path) -> Path:
        """Write a valid PowerPoint file and return its destination."""
        report.validate(artifact_root)
        presentation = Presentation()
        presentation.slide_width = self._slide_width
        presentation.slide_height = self._slide_height
        presentation.core_properties.subject = (
            f"{self.theme.identity}@{self.theme.revision}; "
            f"brand@{self.theme.brand_revision}"
        )

        self._add_title_slide(presentation, report)
        for section_index, section in enumerate(report.sections):
            self._add_section_slide(presentation, report, section, section_index)
            self._add_section_content(
                presentation, report, section, section_index, artifact_root
            )
        if report.findings:
            self._add_summary_slide(presentation, report)
            for finding_index, finding in enumerate(report.findings):
                if finding.visual is not None:
                    self._add_finding_visual(
                        presentation,
                        report,
                        finding,
                        f"findings[{finding_index}].visual",
                        artifact_root,
                    )

        for appendix in report.appendices:
            section = Section(appendix.title, appendix.content, appendix.renderer_hints)
            self._add_section_slide(presentation, report, section, 0)
            self._add_section_content(presentation, report, section, 0, artifact_root)
        destination.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(destination))
        return destination

    def _add_title_slide(self, presentation: Any, report: Report) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_text(
            slide,
            report.metadata.title,
            self._margin,
            Inches(2.1),
            self._content_width,
            Inches(1),
            self.theme.typography.title,
        )
        if report.metadata.subtitle:
            self._add_text(
                slide,
                report.metadata.subtitle,
                self._margin,
                Inches(3.25),
                self._content_width,
                Inches(0.55),
                18,
            )
        if report.metadata.source:
            self._add_text(
                slide,
                report.metadata.source,
                self._margin,
                Inches(6.65),
                self._content_width,
                Inches(0.3),
                10,
                color=self._muted,
            )
        self._add_brand_footer(slide, report, len(presentation.slides))

    def _add_section_slide(
        self,
        presentation: Any,
        report: Report,
        section: Section,
        _section_index: int,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_text(
            slide,
            report.metadata.title,
            self._margin,
            Inches(0.55),
            Inches(10),
            Inches(0.3),
            10,
            color=self._muted,
        )
        self._add_text(
            slide,
            section.title,
            self._margin,
            Inches(3.05),
            self._content_width,
            Inches(0.8),
            28,
        )
        self._add_brand_footer(slide, report, len(presentation.slides))

    def _add_section_content(
        self,
        presentation: Any,
        report: Report,
        section: Section,
        section_index: int,
        artifact_root: Path,
    ) -> None:
        for content_index, content in enumerate(section.content):
            element = f"sections[{section_index}].content[{content_index}]"
            if isinstance(content, ChartReference):
                self._add_chart_slide(
                    presentation, report, section, content, element, artifact_root
                )
            elif isinstance(content, PublicationTable):
                self._add_table_slide(presentation, report, section, content, element)
            else:
                self._add_narrative_slide(presentation, report, section, content)

    def _add_chart_slide(  # noqa: PLR0913, PLR0917
        self,
        presentation: Any,
        report: Report,
        section: Section,
        chart: ChartReference,
        element: str,
        artifact_root: Path,
    ) -> None:
        if chart.artifact is None:
            raise PowerPointRenderError(element, "requires a resolved chart artifact")
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        full_frame = chart.renderer_hints.get("powerpoint.layout") == "full-frame-chart"
        artifact = resolve_chart_artifact(chart, artifact_root)
        if full_frame:
            slide.shapes.add_picture(
                str(artifact), 0, 0, self._slide_width, self._slide_height
            )
            return
        self._add_chrome(
            slide,
            report,
            chart.caption or section.title,
            len(presentation.slides),
        )
        slide.shapes.add_picture(
            str(artifact),
            self._margin,
            self._content_top,
            Inches(8.15),
            self._content_height,
        )
        if chart.caption:
            self._add_text(
                slide,
                chart.caption,
                self._margin,
                Inches(6.75),
                self._content_width,
                Inches(0.3),
                self.theme.typography.caption,
                color=self._muted,
            )

    def _add_table_slide(  # noqa: PLR0913
        self,
        presentation: Any,
        report: Report,
        section: Section,
        table: PublicationTable,
        element: str,
        *,
        finding: Finding | None = None,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_chrome(
            slide,
            report,
            finding.title if finding else table.caption or section.title,
            len(presentation.slides),
            subtitle=finding.subtitle if finding else None,
            source=finding.source if finding else None,
        )
        max_rows = self._hint_integer(
            table.renderer_hints,
            "powerpoint.max_rows",
            _DEFAULT_TABLE_MAX_ROWS,
            element,
        )
        if len(table.rows) > max_rows:
            raise PowerPointRenderError(
                element,
                f"has {len(table.rows)} rows; maximum is {max_rows}",
            )
        left = self._hint_inches(
            table.renderer_hints,
            "powerpoint.left",
            self.theme.geometry.slide_margin,
            element,
        )
        top = self._hint_inches(
            table.renderer_hints,
            "powerpoint.top",
            self.theme.geometry.content_top,
            element,
        )
        width = self._hint_inches(
            table.renderer_hints,
            "powerpoint.width",
            self.theme.geometry.slide_width - 2 * self.theme.geometry.slide_margin,
            element,
        )
        height = self._hint_inches(
            table.renderer_hints, "powerpoint.height", 4.5, element
        )
        font_size = self._hint_integer(
            table.renderer_hints,
            "powerpoint.font_size",
            self.theme.typography.table,
            element,
        )
        row_height = self._hint_inches(
            table.renderer_hints,
            "powerpoint.row_height",
            self.theme.spacing.table_row_inches,
            element,
        )
        rows = len(table.rows) + 1
        shape = slide.shapes.add_table(
            rows, len(table.columns), left, top, width, height
        )
        ppt_table = shape.table
        for column_index, column in enumerate(table.columns):
            ppt_table.cell(0, column_index).text = column.label
            ppt_table.cell(0, column_index).text_frame.paragraphs[0].font.bold = (
                table.renderer_hints.get("powerpoint.header_emphasis", "bold") == "bold"
            )
            column_width = column.renderer_hints.get("powerpoint.width")
            if column_width is not None:
                ppt_table.columns[column_index].width = self._inches(
                    column_width, f"{element}.columns[{column_index}].powerpoint.width"
                )
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row):
                table_cell = ppt_table.cell(row_index, column_index)
                table_cell.text = self._format_cell(
                    cell, table.columns[column_index].format, element
                )
                table_cell.text_frame.paragraphs[0].alignment = self._alignment(
                    table.columns[column_index].alignment
                )
                table_cell.text_frame.paragraphs[0].font.size = Pt(font_size)
        for row in ppt_table.rows:
            row.height = row_height
        for row_index, row in enumerate(ppt_table.rows):
            for cell in row.cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = self._color(
                    self.theme.colors.primary
                    if row_index == 0
                    else self.theme.colors.background
                )
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = self.theme.typography.family
                    paragraph.font.size = Pt(font_size)
                    paragraph.font.color.rgb = self._color(
                        self.theme.colors.background
                        if row_index == 0
                        else self.theme.colors.text_primary
                    )
        if table.caption:
            self._add_text(
                slide,
                table.caption,
                self._margin,
                Inches(6.15),
                Inches(12),
                Inches(0.35),
                self.theme.typography.caption,
                color=self._muted,
            )
        if table.notes:
            self._add_text(
                slide,
                " ".join(table.notes),
                self._margin,
                Inches(6.55),
                Inches(12),
                Inches(0.3),
                self.theme.typography.note,
                color=self._muted,
            )

    def _add_narrative_slide(
        self, presentation: Any, report: Report, section: Section, narrative: Narrative
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_chrome(slide, report, section.title, len(presentation.slides))
        style = self.theme.primitive(narrative.kind, narrative.semantic)
        text = narrative.text
        if style.label:
            self._add_text(
                slide,
                style.label,
                self._margin,
                self._content_top,
                self._content_width,
                Inches(0.4),
                self.theme.typography.heading,
                color=self._color(style.color),
                bold=True,
            )
        shape = self._add_text(
            slide,
            text,
            self._margin,
            self._content_top + (Inches(0.55) if style.label else 0),
            self._content_width,
            Inches(4.6),
            self.theme.typography.heading
            if narrative.kind == "heading"
            else self.theme.typography.body,
            color=self._color(style.color),
            bold=style.bold and narrative.kind == "heading",
            markdown=narrative.kind != "code",
            font=self.theme.typography.code_family
            if narrative.kind == "code"
            else None,
        )
        if narrative.kind == "callout":
            shape.line.color.rgb = self._color(style.color)
            shape.line.width = Pt(2)
        if narrative.kind == "list":
            frame = shape.text_frame
            frame.clear()
            for index, entry in enumerate(list_entries(text)):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.level = min(entry.level, 8)
                self._paragraph(
                    paragraph,
                    f"{entry.marker} {entry.text}".strip(),
                    self.theme.typography.body,
                    self._color(style.color),
                    markdown=True,
                )

    def _add_summary_slide(self, presentation: Any, report: Report) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_chrome(slide, report, "Summary", len(presentation.slides))
        for index, finding in enumerate(report.findings):
            self._add_finding(slide, finding, index)

    def _add_finding_visual(
        self,
        presentation: Any,
        report: Report,
        finding: Finding,
        element: str,
        artifact_root: Path,
    ) -> None:
        visual = finding.visual
        if visual is None:
            raise PowerPointRenderError(element, "requires a finding visual")
        if isinstance(visual, ChartReference):
            if visual.artifact is None:
                raise PowerPointRenderError(
                    element, "requires a resolved chart artifact"
                )
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            if visual.renderer_hints.get("powerpoint.layout") == "full-frame-chart":
                slide.shapes.add_picture(
                    str(resolve_chart_artifact(visual, artifact_root)),
                    0,
                    0,
                    self._slide_width,
                    self._slide_height,
                )
                return
            self._add_chrome(
                slide,
                report,
                finding.title,
                len(presentation.slides),
                subtitle=finding.subtitle,
                source=finding.source,
            )
            # Finding chart references may be absolute or already materialized.
            slide.shapes.add_picture(
                str(resolve_chart_artifact(visual, artifact_root)),
                self._margin,
                self._content_top,
                Inches(8.15),
                self._content_height,
            )
        else:
            # Table semantics and constraints are shared with report-native tables.
            self._add_table_slide(
                presentation,
                report,
                Section(finding.title, (visual,)),
                visual,
                element,
                finding=finding,
            )

    def _add_finding(self, slide: Any, finding: Finding, index: int) -> None:
        top = self._content_top + Inches(index * 1.25)
        self._add_text(
            slide,
            finding.title,
            self._margin,
            top,
            self._content_width,
            Inches(0.3),
            16,
            color=self._color(self.theme.finding_color(finding.severity)),
            bold=True,
        )
        self._add_text(
            slide,
            finding.narrative.text,
            self._margin,
            top + Inches(0.35),
            Inches(12),
            Inches(0.65),
            12,
            markdown=True,
        )

    def _add_chrome(  # noqa: PLR0913
        self,
        slide: Any,
        report: Report,
        heading: str,
        page: int,
        *,
        subtitle: str | None = None,
        source: str | None = None,
    ) -> None:
        self._add_text(
            slide,
            report.metadata.title,
            self._margin,
            Inches(0.35),
            Inches(10),
            Inches(0.3),
            10,
            color=self._muted,
        )
        self._add_text(
            slide,
            heading,
            self._margin,
            Inches(0.75),
            Inches(10),
            Inches(0.55),
            self.theme.typography.heading,
        )
        if subtitle or report.metadata.subtitle:
            self._add_text(
                slide,
                subtitle or report.metadata.subtitle or "",
                self._margin,
                Inches(1.18),
                Inches(10),
                Inches(0.25),
                9,
                color=self._muted,
            )
        if source or report.metadata.source:
            self._add_text(
                slide,
                source or report.metadata.source or "",
                self._margin,
                Inches(7.1),
                Inches(8),
                Inches(0.2),
                8,
                color=self._muted,
            )
        self._add_brand_footer(slide, report, page)

    def _add_brand_footer(self, slide: Any, report: Report, page: int) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self._color(self.theme.colors.background)
        self._add_text(
            slide,
            self.theme.wordmark,
            self._slide_width - Inches(2.8),
            Inches(0.35),
            Inches(2.2),
            Inches(0.3),
            11,
            bold=True,
        )
        if report.metadata.confidentiality:
            self._add_text(
                slide,
                report.metadata.confidentiality,
                self._slide_width - Inches(4.8),
                self._slide_height - Inches(0.5),
                Inches(3.4),
                Inches(0.25),
                self.theme.typography.note,
                color=self._muted,
                alignment=PP_ALIGN.RIGHT,
            )
        if self.theme.page_numbers:
            self._add_page_number(slide, page)

    @staticmethod
    def _format_cell(cell: object, format_spec: str | None, element: str) -> str:
        if cell is None:
            return ""
        try:
            return format(cell, format_spec or "")
        except (TypeError, ValueError) as error:
            raise PowerPointRenderError(
                element, f"invalid table format {format_spec!r}: {error}"
            ) from error

    def _add_page_number(self, slide: Any, page: int) -> None:
        self._add_text(
            slide,
            str(page),
            self._slide_width - Inches(1.2),
            self._slide_height - Inches(0.5),
            Inches(0.5),
            Inches(0.2),
            8,
            color=self._muted,
            alignment=PP_ALIGN.RIGHT,
        )

    @staticmethod
    def _alignment(alignment: str) -> PP_ALIGN:
        return {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[alignment]

    def _hint_integer(
        self, hints: dict[str, str], name: str, default: int, element: str
    ) -> int:
        value = hints.get(name)
        if value is None:
            return default
        try:
            result = int(value)
        except ValueError as error:
            raise PowerPointRenderError(
                element, f"{name} must be an integer"
            ) from error
        if result <= 0:
            raise PowerPointRenderError(element, f"{name} must be positive")
        return result

    def _hint_inches(
        self, hints: dict[str, str], name: str, default: float, element: str
    ) -> Any:
        value = hints.get(name)
        return (
            Inches(default)
            if value is None
            else self._inches(value, f"{element}.{name}")
        )

    @staticmethod
    def _inches(value: str, element: str) -> Any:
        try:
            result = float(value)
        except ValueError as error:
            raise PowerPointRenderError(
                element, "must be a number of inches"
            ) from error
        if result <= 0:
            raise PowerPointRenderError(element, "must be positive")
        return Inches(result)

    def _add_text(  # noqa: PLR0913, PLR0917
        self,
        slide: Any,
        text: str,
        left: Any,
        top: Any,
        width: Any,
        height: Any,
        size: Any,
        *,
        color: Any = None,
        alignment: PP_ALIGN = PP_ALIGN.LEFT,
        bold: bool = False,
        markdown: bool = False,
        font: str | None = None,
    ) -> Any:
        shape = slide.shapes.add_textbox(left, top, width, height)
        frame = shape.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.alignment = alignment
        self._paragraph(
            paragraph,
            text,
            size,
            color or self._primary,
            bold=bold,
            markdown=markdown,
            font=font,
        )
        return shape

    def _paragraph(  # noqa: PLR0913
        self,
        paragraph: Any,
        text: str,
        size: Any,
        color: Any,
        *,
        bold: bool = False,
        markdown: bool = False,
        font: str | None = None,
    ) -> None:
        paragraph.font.name = font or self.theme.typography.family
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.font.bold = bold
        paragraph.space_after = Pt(self.theme.spacing.paragraph_points)
        if markdown:
            for item in text_runs(text):
                run = paragraph.add_run()
                run.text = item.text
                run.font.bold = bold or item.bold
                run.font.italic = item.italic
                if item.code:
                    run.font.name = self.theme.typography.code_family
                if item.hyperlink:
                    run.hyperlink.address = item.hyperlink
        else:
            paragraph.text = text
