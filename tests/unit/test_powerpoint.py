import base64
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from wood_reports import (
    ChartReference,
    Finding,
    Narrative,
    PowerPointRenderer,
    PowerPointRenderError,
    PublicationTable,
    Report,
    ReportMetadata,
    Section,
    TableColumn,
)


def write_png(path: Path) -> None:
    path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
            "OcE1owAAAABJRU5ErkJggg=="
        )
    )


def report_with_layouts(chart: Path) -> Report:
    return Report(
        metadata=ReportMetadata(
            title="Monthly performance",
            subtitle="August 2026",
            source="Finance ledger",
        ),
        sections=(
            Section(
                "Overview",
                (
                    Narrative("Revenue increased during the month."),
                    ChartReference(artifact=chart, caption="Revenue trend"),
                    PublicationTable(
                        columns=(
                            TableColumn("month", "Month"),
                            TableColumn("revenue", "Revenue"),
                        ),
                        rows=(("August", 1200),),
                        caption="Monthly revenue",
                    ),
                ),
            ),
        ),
        findings=(
            Finding(
                "growth", "Revenue increased", Narrative("Growth was broad based.")
            ),
        ),
    )


def slide_text(slide: Any) -> str:
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def test_renderer_writes_valid_editable_presentation_with_all_layouts(
    tmp_path: Path,
) -> None:
    chart = tmp_path / "charts" / "revenue.png"
    chart.parent.mkdir()
    write_png(chart)
    destination = tmp_path / "reports" / "monthly.pptx"

    output = PowerPointRenderer().render(
        report_with_layouts(chart), destination, artifact_root=tmp_path
    )
    presentation = Presentation(str(output))

    assert output == destination
    assert destination.is_file()
    assert len(presentation.slides) == 6
    assert "Monthly performance" in slide_text(presentation.slides[3])
    assert "August 2026" in slide_text(presentation.slides[3])
    assert "Finance ledger" in slide_text(presentation.slides[3])
    assert "4" in slide_text(presentation.slides[3])
    assert any(shape.shape_type == 13 for shape in presentation.slides[3].shapes)
    assert any(shape.has_table for shape in presentation.slides[4].shapes)
    assert "Revenue increased" in slide_text(presentation.slides[5])


def test_full_frame_chart_layout_uses_the_entire_slide_without_report_chrome(
    tmp_path: Path,
) -> None:
    chart = tmp_path / "chart.png"
    write_png(chart)
    report = Report(
        metadata=ReportMetadata(title="Monthly performance"),
        sections=(
            Section(
                "Overview",
                (
                    ChartReference(
                        artifact=chart,
                        renderer_hints={"powerpoint.layout": "full-frame-chart"},
                    ),
                ),
            ),
        ),
    )
    destination = tmp_path / "full-frame.pptx"

    PowerPointRenderer().render(report, destination, artifact_root=tmp_path)
    presentation = Presentation(str(destination))

    assert len(presentation.slides) == 3
    assert slide_text(presentation.slides[2]) == ""
    assert len(presentation.slides[2].shapes) == 1


def test_table_hints_control_native_layout_without_mutating_source_data(
    tmp_path: Path,
) -> None:
    table = PublicationTable(
        columns=(
            TableColumn("name", "Name", renderer_hints={"powerpoint.width": "3"}),
            TableColumn("value", "Value", alignment="right"),
        ),
        rows=(("August", 1200.5),),
        caption="Monthly revenue",
        notes=("Unaudited.",),
        renderer_hints={
            "powerpoint.left": "1",
            "powerpoint.top": "2",
            "powerpoint.width": "10",
            "powerpoint.height": "3",
            "powerpoint.font_size": "14",
            "powerpoint.row_height": "0.5",
            "powerpoint.max_rows": "2",
        },
    )
    report = Report(
        metadata=ReportMetadata("Table report"),
        sections=(Section("Overview", (table,)),),
    )
    destination = tmp_path / "table.pptx"

    PowerPointRenderer().render(report, destination, artifact_root=tmp_path)
    rendered_table = next(
        shape.table
        for shape in Presentation(str(destination)).slides[2].shapes
        if shape.has_table
    )

    assert rendered_table.cell(1, 0).text == "August"
    assert rendered_table.cell(1, 1).text == "1200.5"
    assert (
        rendered_table.cell(1, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    )
    assert rendered_table.cell(0, 0).text_frame.paragraphs[0].font.bold is True
    assert table.rows == (("August", 1200.5),)


def test_table_row_limit_reports_the_table_location(tmp_path: Path) -> None:
    report = Report(
        metadata=ReportMetadata("Table report"),
        sections=(
            Section(
                "Overview",
                (
                    PublicationTable(
                        columns=(TableColumn("name", "Name"),),
                        rows=(("August",), ("September",)),
                        renderer_hints={"powerpoint.max_rows": "1"},
                    ),
                ),
            ),
        ),
    )

    with pytest.raises(PowerPointRenderError, match=r"sections\[0\]\.content\[0\]"):
        PowerPointRenderer().render(
            report, tmp_path / "too-many-rows.pptx", artifact_root=tmp_path
        )


def test_finding_visual_uses_editable_finding_chrome(tmp_path: Path) -> None:
    chart = tmp_path / "chart.png"
    write_png(chart)
    report = Report(
        ReportMetadata("Report"),
        (Section("Results", (Narrative("Text"),)),),
        findings=(
            Finding(
                "weekday",
                "Weekend falls",
                Narrative("Narrative"),
                subtitle="Below baseline",
                source="Synthetic",
                visual=ChartReference(Path("chart.png")),
            ),
        ),
    )
    output = tmp_path / "finding.pptx"
    PowerPointRenderer().render(report, output, artifact_root=tmp_path)
    text = slide_text(Presentation(str(output)).slides[-1])
    assert "Weekend falls" in text
    assert "Below baseline" in text
    assert "Synthetic" in text
