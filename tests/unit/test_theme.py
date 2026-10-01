import base64
import json
from dataclasses import replace
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import pytest
from pptx import Presentation
from wood_charts import load_theme

from wood_reports import (
    WOOD_ANALYTICS_THEME,
    Appendix,
    ChartReference,
    Finding,
    LatexRenderer,
    MarkdownReportCompiler,
    Narrative,
    PowerPointRenderer,
    PublicationTable,
    Report,
    ReportGenerationAPI,
    ReportMetadata,
    ReportRun,
    Section,
    TableColumn,
    get_profile,
)
from wood_reports.primitives import latex_escape, list_entries, text_runs
from wood_reports.theme import PublicationTheme


def semantic_report() -> Report:
    profile = get_profile("decision-memo")
    text = (
        "---\ndoc_type: decision-memo\ndoc_name: internal-example\n"
        "title: 'Internal & example'\nconfidentiality: Internal\nsource: Ledger\n---\n"
    )
    for section in profile.sections:
        if section.identity in profile.required_section_variants[0]:
            text += (
                f"\n## {section.title}\n\nPlain **strong** and *emphasis*, `code`. "
                "[Evidence](https://example.com/evidence?a=1&b=2).\n"
            )
    text += """

## Appendix

### Details

- First
  - Nested
- Last

3. Third
4. Fourth

> [!RISK] Assumptions & uncertainty.

```python
value = "50%"
```

| Measure | Value |
| --- | ---: |
| A & B | 50% |
"""
    return MarkdownReportCompiler().compile_text(text)


def all_text(slide: Any) -> str:
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def test_shared_tokens_match_the_pinned_chart_package() -> None:
    chart_theme = load_theme("base")
    theme = WOOD_ANALYTICS_THEME
    for name in (
        "primary",
        "secondary",
        "text_primary",
        "text_secondary",
        "text_muted",
        "neutral_light",
        "grid",
        "background",
    ):
        assert getattr(theme.colors, name) == getattr(chart_theme.colors, name)
    assert chart_theme.typography.family.split(", ")[0] == theme.typography.family
    assert theme.typography.fallback in chart_theme.typography.family
    # The vector asset and editable wordmark equivalents share tokens.
    asset = ElementTree.fromstring(theme.wordmark_svg)  # noqa: S314
    wordmark = asset.find("{http://www.w3.org/2000/svg}text")
    assert wordmark is not None
    assert wordmark.text == theme.wordmark
    assert wordmark.attrib["fill"] == theme.colors.primary


@pytest.mark.parametrize("semantic", get_profile("decision-memo").callouts)
def test_every_profile_callout_has_a_deterministic_style(semantic: str) -> None:
    style = WOOD_ANALYTICS_THEME.primitive("callout", semantic)
    assert style.label == semantic.title()
    assert style.bold
    assert style.color == (
        WOOD_ANALYTICS_THEME.colors.warning
        if semantic == "risk"
        else WOOD_ANALYTICS_THEME.colors.primary
    )


@pytest.mark.parametrize(
    ("kind", "semantic"), [("callout", "unknown"), ("unknown", None)]
)
def test_unknown_intent_is_not_silently_flattened(
    kind: str, semantic: str | None
) -> None:
    with pytest.raises(ValueError, match="unsupported"):
        WOOD_ANALYTICS_THEME.primitive(kind, semantic)


def test_markdown_semantics_render_as_branded_latex(tmp_path: Path) -> None:
    report = semantic_report()
    output = LatexRenderer().render(
        report, tmp_path / "report.tex", artifact_root=tmp_path
    )
    text = output.read_text()
    assert r"\definecolor{Woodprimary}{HTML}{002F6C}" in text
    assert r"\pagecolor{Woodbackground}" in text
    assert r"\setsansfont{lmsans10-regular.otf}" in text
    assert r"\setmonofont{lmmono10-regular.otf}" in text
    assert r"\geometry{paperwidth=8.5in,paperheight=11.0in,margin=0.8in" in text
    assert r"\title{Internal \& example}" in text
    assert r"\textbf{strong}" in text
    assert r"\emph{emphasis}" in text
    assert r"\texttt{code}" in text
    assert r"\appendix" in text
    assert r"\subsection*{Details}" in text
    assert r"\item[{3.}]" in text
    assert r"\item[{\textbullet}] Nested" in text
    assert text.count(r"\begin{list}") == 3
    assert r"\href{https://example.com/evidence?a=1\&b=2}{Evidence}" in text
    assert r"\WoodCallout{Woodwarning}{Risk}{Assumptions \& uncertainty.}" in text
    assert r"\ttfamily value\ =\ " in text
    assert r"\rowcolor{Woodprimary}" in text
    assert r"A \& B" in text
    assert r"50\%" in text
    assert r"\WoodConfidentiality{Internal}" in text
    assert r"\WoodSource{Ledger}" in text


def test_markdown_semantics_remain_editable_in_powerpoint(tmp_path: Path) -> None:
    report = semantic_report()
    output = PowerPointRenderer().render(
        report, tmp_path / "report.pptx", artifact_root=tmp_path
    )
    deck = Presentation(str(output))
    texts = [all_text(slide) for slide in deck.slides]
    assert all("Wood Analytics" in text and "Internal" in text for text in texts)
    assert any(
        "Risk" in text and "Assumptions & uncertainty." in text for text in texts
    )
    assert any("3. Third" in text and "4. Fourth" in text for text in texts)
    body_slide = next(
        slide for slide in deck.slides if "Plain strong" in all_text(slide)
    )
    body = next(
        shape
        for shape in body_slide.shapes
        if shape.has_text_frame and "Plain strong" in shape.text
    )
    paragraph = body.text_frame.paragraphs[0]
    assert paragraph.font.name == "IBM Plex Sans"
    assert any(run.text == "strong" and run.font.bold for run in paragraph.runs)
    assert any(run.text == "emphasis" and run.font.italic for run in paragraph.runs)
    assert any(
        run.text == "code" and run.font.name == "IBM Plex Mono"
        for run in paragraph.runs
    )
    assert any(
        run.text == "Evidence"
        and run.hyperlink.address == "https://example.com/evidence?a=1&b=2"
        for run in paragraph.runs
    )
    nested_slide = next(slide for slide in deck.slides if "• First" in all_text(slide))
    nested = next(
        shape
        for shape in nested_slide.shapes
        if shape.has_text_frame and "Nested" in shape.text
    )
    assert [p.level for p in nested.text_frame.paragraphs] == [0, 1, 0]
    table = next(
        shape.table
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_table
    )
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "002F6C"
    assert str(table.cell(0, 0).text_frame.paragraphs[0].font.color.rgb) == "FFFFFF"
    assert str(table.cell(1, 0).text_frame.paragraphs[0].font.color.rgb) == "1B1B1B"
    assert deck.core_properties.subject == "wood-analytics@1.0.0; brand@1.0.0"


def test_custom_theme_survives_materialization_and_both_manifest_targets(
    tmp_path: Path,
) -> None:
    original = semantic_report()
    theme = replace(
        original.theme,
        revision="1.1.0",
        brand_revision="1.1.0",
        page_numbers=False,
        colors=replace(original.theme.colors, primary="#123456"),
        geometry=replace(original.theme.geometry, slide_width=14),
    )
    report = replace(original, theme=theme)
    result = ReportGenerationAPI().generate(
        ReportRun.create(report, period="2026-10"),
        tmp_path / "output",
        artifact_root=tmp_path,
    )
    assert all(target.status == "success" for target in result.targets)
    assert result.report.theme is theme
    assert result.manifest is not None
    assert json.loads(result.manifest.read_text())["theme"] == {
        "identity": "wood-analytics",
        "revision": "1.1.0",
        "brand_revision": "1.1.0",
    }
    assert (
        r"\definecolor{Woodprimary}{HTML}{123456}"
        in (result.manifest.parent / "report.tex").read_text()
    )
    deck = Presentation(str(result.manifest.parent / "report.pptx"))
    assert deck.slide_width == 14 * 914400
    assert not any(
        shape.text == "1" for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    assert original.theme is WOOD_ANALYTICS_THEME


def test_findings_figures_tables_and_notes_share_the_theme(tmp_path: Path) -> None:
    chart = tmp_path / "chart.png"
    chart.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
            "OcE1owAAAABJRU5ErkJggg=="
        )
    )
    table = PublicationTable(
        (TableColumn("x", "X"),), ((1,),), caption="Caption", notes=("Source note.",)
    )
    report = Report(
        ReportMetadata("Example"),
        (Section("Data", (ChartReference(chart, caption="Figure caption"),)),),
        findings=(
            Finding(
                "risk",
                "Risk finding",
                Narrative("Watch."),
                severity="critical",
                visual=table,
            ),
        ),
        appendices=(Appendix("Extra", (Narrative("Tail."),)),),
    )
    output = PowerPointRenderer().render(
        report, tmp_path / "report.pptx", artifact_root=tmp_path
    )
    deck = Presentation(str(output))
    assert any("Figure caption" in all_text(slide) for slide in deck.slides)
    assert any("Source note." in all_text(slide) for slide in deck.slides)
    summary = next(slide for slide in deck.slides if "Summary" in all_text(slide))
    title = next(
        shape
        for shape in summary.shapes
        if shape.has_text_frame and shape.text == "Risk finding"
    )
    assert str(title.text_frame.paragraphs[0].font.color.rgb) == "991B1B"
    text = (
        LatexRenderer()
        .render(report, tmp_path / "report.tex", artifact_root=tmp_path)
        .read_text()
    )
    assert r"\WoodCallout{Woodcritical}{Critical}{Watch.}" in text
    assert "Source note." in text


def test_inline_and_list_interpretation_preserves_special_text() -> None:
    assert (
        latex_escape("&%$#_{}~^\\")
        == r"\&\%\$\#\_\{\}\textasciitilde{}\textasciicircum{}\textbackslash{}"
    )
    assert [run.text for run in text_runs("one\ntwo  \nthree")] == [
        "one",
        "\n",
        "two",
        "\n",
        "three",
    ]
    entries = list_entries("- First\n\n  Continuation\n\n- Last")
    assert [entry.marker for entry in entries] == ["•", "", "•"]
    assert [entry.text for entry in entries] == ["First", "Continuation", "Last"]


@pytest.mark.parametrize(
    "theme",
    [
        replace(WOOD_ANALYTICS_THEME, identity=""),
        replace(
            WOOD_ANALYTICS_THEME,
            colors=replace(WOOD_ANALYTICS_THEME.colors, primary="red"),
        ),
        replace(
            WOOD_ANALYTICS_THEME,
            geometry=replace(WOOD_ANALYTICS_THEME.geometry, slide_width=float("inf")),
        ),
        replace(
            WOOD_ANALYTICS_THEME,
            typography=replace(WOOD_ANALYTICS_THEME.typography, family=""),
        ),
        replace(
            WOOD_ANALYTICS_THEME,
            typography=replace(WOOD_ANALYTICS_THEME.typography, body=0),
        ),
        replace(
            WOOD_ANALYTICS_THEME,
            geometry=replace(WOOD_ANALYTICS_THEME.geometry, slide_margin=8),
        ),
    ],
)
def test_invalid_theme_is_rejected_before_output(
    theme: PublicationTheme, tmp_path: Path
) -> None:
    report = replace(semantic_report(), theme=theme)
    with pytest.raises(ValueError, match="theme"):
        LatexRenderer().render(report, tmp_path / "report.tex", artifact_root=tmp_path)
    assert not (tmp_path / "report.tex").exists()
