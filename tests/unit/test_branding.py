"""Cross-profile contract tests against actual generated documents."""

import json
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches

from wood_reports import (
    LatexRenderer,
    LatexWorkspacePublisher,
    MarkdownReportCompiler,
    PowerPointRenderer,
    PublicationBranding,
    Report,
    list_profiles,
    scaffold_markdown,
)
from wood_reports.branding import parse_logo, require_fonts
from wood_reports.clsi_workspace import read_workspace
from wood_reports.theme import PublicationTheme

REPLACEMENT = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 20">'
    '<path fill="#123456" d="M0 0 H100 V20 H0 Z M20 5 V15 H80 V5 Z"/></svg>'
)


def profile_report(identity: str) -> Report:
    text = scaffold_markdown(identity, "internal-branding")
    text = text.replace(
        "title: Internal Branding",
        "title: Internal Branding\nconfidentiality: Internal",
    )
    text = "\n".join(
        line + "\n\nInternal evidence." if line.startswith("## ") else line
        for line in text.splitlines()
    )
    text += (
        "\n\n## Appendix\n\n```python\nvalue = 42\n```\n\n"
        "| Measure | Value |\n| --- | ---: |\n| Coverage | 42 |\n"
    )
    return MarkdownReportCompiler().compile_text(text)


@pytest.mark.parametrize("profile", [profile.identity for profile in list_profiles()])
@pytest.mark.parametrize(
    "corner", ["top-left", "top-right", "bottom-left", "bottom-right"]
)
@pytest.mark.parametrize("custom", [False, True])
def test_all_profiles_share_portable_vector_branding(
    profile: str, corner: Any, custom: bool, tmp_path: Path
) -> None:
    report = profile_report(profile)
    branding = PublicationBranding(
        corner=corner, logo_svg=REPLACEMENT if custom else None
    )
    report = replace(report, theme=replace(report.theme, branding=branding))
    tex = LatexWorkspacePublisher().publish(
        report, tmp_path / "workspace", artifact_root=tmp_path
    )
    manifest = json.loads((tex.parent / "workspace.json").read_text())
    assert manifest["profile"] == {"identity": profile, "version": "1.0.0"}
    assert manifest["brand"]["identity"] == "wood-analytics"
    assert (tex.parent / "assets/brand.svg").read_text() == report.theme.wordmark_svg
    assert r"\input{components.tex}" in tex.read_text()
    components = (tex.parent / "components.tex").read_text()
    assert "WoodBrandLogo" in components
    assert "IBM Plex Sans" in components
    assert "IBM Plex Mono" in components
    assert "WoodTableFont" in components
    assert "woodcaption" in components
    assert "Internal" in components
    assert r"\thepage" in components
    inputs = read_workspace(tex.parent.parent, {})
    assert "generated/components.tex" in inputs.fingerprints
    assert "generated/assets/brand.svg" in inputs.fingerprints
    deck = Presentation(
        str(
            PowerPointRenderer().render(
                report, tmp_path / "report.pptx", artifact_root=tmp_path
            )
        )
    )
    x, y = branding.position(
        report.theme.geometry.slide_width, report.theme.geometry.slide_height
    )
    for slide in deck.slides:
        logos = [shape for shape in slide.shapes if shape.name.startswith("Wood brand")]
        assert logos
        if custom:
            shape = logos[0]
            assert abs(shape.left - Inches(x)) < 10
            assert abs(shape.top - Inches(y)) < 10
            expected_width = min(branding.width, branding.height * 5)
            assert abs(shape.width - Inches(expected_width)) < 100
            assert str(shape.fill.fore_color.rgb) == "123456"
        else:
            assert any(
                shape.has_text_frame and shape.text == "Wood Analytics"
                for shape in logos
            )


@pytest.mark.parametrize(
    ("visible", "cover"), [(False, False), (False, True), (True, False), (True, True)]
)
def test_cover_and_body_visibility_are_independent(
    visible: bool, cover: bool, tmp_path: Path
) -> None:
    report = profile_report("decision-memo")
    report = replace(
        report,
        theme=replace(
            report.theme,
            branding=PublicationBranding(visible=visible, cover_visible=cover),
        ),
    )
    deck = Presentation(
        str(
            PowerPointRenderer().render(
                report, tmp_path / "report.pptx", artifact_root=tmp_path
            )
        )
    )
    for index, slide in enumerate(deck.slides):
        assert any(shape.name.startswith("Wood brand") for shape in slide.shapes) == (
            cover if index == 0 else visible
        )
        assert any(
            shape.has_text_frame and shape.text == "Internal" for shape in slide.shapes
        )
    text = (
        LatexRenderer()
        .render(report, tmp_path / "report.tex", artifact_root=tmp_path)
        .read_text()
    )
    overlay = text.split(r"\AddToShipoutPictureFG", 1)[1].split("\n", 1)[0]
    first, body = overlay.split(r"\else", 1)
    assert ("WoodBrandLogo" in first) == cover
    assert ("WoodBrandLogo" in body) == visible


@pytest.mark.parametrize(
    "options",
    [
        {"corner": "middle"},
        {"width": 0},
        {"width": float("nan")},
        {"height": 100},
        {"offset_x": -1},
        {"offset_y": float("inf")},
        {"visible": "yes"},
        {"font_policy": "implicit"},
    ],
)
def test_invalid_configuration_fails_before_outputs(
    options: dict[str, Any], tmp_path: Path
) -> None:
    report = profile_report("project-brief")
    report = replace(
        report, theme=replace(report.theme, branding=PublicationBranding(**options))
    )
    with pytest.raises(ValueError, match="branding"):
        LatexRenderer().render(report, tmp_path / "report.tex", artifact_root=tmp_path)
    assert not (tmp_path / "report.tex").exists()


@pytest.mark.parametrize(
    "svg",
    [
        "not XML",
        "<!DOCTYPE svg><svg/>",
        '<svg xmlns="http://www.w3.org/2000/svg"/>',
        REPLACEMENT.replace('fill="#123456"', 'fill="url(https://example.com)"'),
        REPLACEMENT.replace("<path", '<path transform="rotate(10)"'),
        REPLACEMENT.replace('d="M0 0 H100 V20 H0 Z M20 5 V15 H80 V5 Z"', 'd="M0 0 L"'),
        REPLACEMENT.replace("<path", "<script"),
        REPLACEMENT.replace("#123456", "red"),
    ],
)
def test_unsupported_and_malformed_svg_are_actionable(svg: str) -> None:
    with pytest.raises(ValueError, match=r"branding\.logo_svg"):
        parse_logo(svg)


def test_asset_snapshot_is_portable_and_missing_assets_fail(tmp_path: Path) -> None:
    asset = tmp_path / "logo.svg"
    asset.write_text(REPLACEMENT)
    branding = PublicationBranding.from_svg(asset)
    asset.unlink()
    assert branding.logo_svg == REPLACEMENT
    with pytest.raises(ValueError, match="cannot read"):
        PublicationBranding.from_svg(asset)


def test_curves_arcs_and_multiple_contours_are_supported() -> None:
    svg = REPLACEMENT.replace(
        "M0 0 H100 V20 H0 Z M20 5 V15 H80 V5 Z",
        "M0 0 C20 0 30 20 40 20 Q50 0 60 20 A10 10 0 0 1 80 20 L100 0 Z",
    )
    logo = parse_logo(svg)
    assert len(logo.shapes[0].contours[0]) > 100


def test_brand_and_profile_revisions_change_independently(tmp_path: Path) -> None:
    report = profile_report("decision-memo")
    brand_change = replace(
        report,
        theme=replace(
            report.theme, brand_identity="replacement", brand_revision="2.0.0"
        ),
    )
    profile_change = replace(
        report, metadata=replace(report.metadata, profile_version="2.0.0")
    )
    manifests = []
    for index, value in enumerate((report, brand_change, profile_change)):
        output = LatexWorkspacePublisher().publish(
            value, tmp_path / str(index), artifact_root=tmp_path
        )
        manifests.append(json.loads((output.parent / "workspace.json").read_text()))
    assert manifests[0]["profile"] == manifests[1]["profile"]
    assert manifests[0]["brand"] == manifests[2]["brand"]
    assert manifests[1]["brand"]["revision"] == "2.0.0"
    assert manifests[2]["profile"]["version"] == "2.0.0"


def test_strict_fonts_reject_substitution(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("wood_reports.branding.shutil.which", lambda _: None)
    with pytest.raises(ValueError, match="fontconfig"):
        require_fonts(("Missing family",))


def test_publication_config_snapshots_one_central_brand(tmp_path: Path) -> None:
    (tmp_path / "logo.svg").write_text(REPLACEMENT)
    config = tmp_path / "pyproject.toml"
    config.write_text(
        '[tool.wood_reports.publication]\nbrand_identity = "internal"\n'
        'brand_revision = "2.0.0"\n[tool.wood_reports.publication.branding]\n'
        'svg_path = "logo.svg"\ncorner = "bottom-left"\ncover_visible = false\n'
        "[tool.wood_reports.publication.typography]\ncaption = 12\n"
    )
    theme = PublicationTheme.from_pyproject(config)
    assert theme.brand_identity == "internal"
    assert theme.branding.logo_svg == REPLACEMENT
    assert theme.branding.cover_visible is False
    assert theme.typography.caption == 12
    (tmp_path / "logo.svg").unlink()
    theme.validate()
    with pytest.raises(ValueError, match="cannot read"):
        PublicationTheme.from_pyproject(config)


@pytest.mark.parametrize(
    "config",
    [
        "[tool.wood_reports.publication]\nunknown = true",
        '[tool.wood_reports.publication.branding]\nwidth = "large"',
        '[tool.wood_reports.publication.branding]\nsvg_path = "x"\nlogo_svg = "x"',
        '[tool.wood_reports.publication.colors]\nprimary = "red"',
    ],
)
def test_unsupported_configuration_is_rejected(config: str, tmp_path: Path) -> None:
    path = tmp_path / "pyproject.toml"
    path.write_text(config)
    with pytest.raises(ValueError, match=r"tool\.wood_reports\.publication"):
        PublicationTheme.from_pyproject(path)


def test_custom_wordmark_and_palette_share_the_placeholder_asset() -> None:
    theme = PublicationTheme(
        wordmark="Internal & Co",
        colors=replace(PublicationTheme().colors, primary="#123456"),
    )
    logo = parse_logo(theme.wordmark_svg)
    assert logo.shapes[1].text == "Internal & Co"
    assert all(shape.color == "#123456" for shape in logo.shapes)
